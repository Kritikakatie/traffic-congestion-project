"""REPORT.md -> REPORT.docx (hard-copy friendly: page breaks, 6-inch figures, bordered tables). Needs pandoc and python-docx."""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
text = (ROOT / "REPORT.md").read_text(encoding="utf-8")
text = text.replace("\\newpage", '```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```')
text = re.sub(r"(!\[[^\]]*\]\([^)]+\))", r"\1{width=6in}", text)
with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8", dir=ROOT) as fh:
    fh.write(text)
    tmp = Path(fh.name)
out = ROOT / "REPORT.docx"
try:
    subprocess.run(["pandoc", str(tmp), "-o", str(out), "--resource-path", str(ROOT), "-f", "markdown+pipe_tables"], check=True, cwd=ROOT)
finally:
    tmp.unlink(missing_ok=True)

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

doc = Document(str(out))
for t in doc.tables:
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "808080")
        borders.append(el)
    tblPr.append(borders)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
doc.save(str(out))
print("wrote", out)
